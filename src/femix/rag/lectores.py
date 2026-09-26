"""Sacar el texto de lo que sube un inquilino: PDF, Word, Excel, HTML, CSV, Markdown o texto.

Cada lector devuelve texto plano con los títulos marcados con `# ` cuando se pueden saber (Word,
HTML, hojas de Excel), para que `fragmentos.fragmentar_por_secciones` corte por apartados. Un
fichero ilegible o de un tipo que no se conoce da ValueError con un mensaje para el usuario.
"""
import csv
import io
import os
from html.parser import HTMLParser

EXTENSIONES = (".txt", ".md", ".csv", ".pdf", ".docx", ".xlsx", ".html", ".htm")
MAXIMO_CARACTERES = 2_000_000   # lo que se indexa como mucho de un solo fichero


def _texto_plano(contenido: bytes) -> str:
    for codificacion in ("utf-8-sig", "latin-1"):
        try:
            return contenido.decode(codificacion)
        except UnicodeDecodeError:
            continue
    return contenido.decode("utf-8", errors="ignore")


def _pdf(contenido: bytes) -> str:
    from pypdf import PdfReader
    lector = PdfReader(io.BytesIO(contenido))
    if lector.is_encrypted:
        try:
            lector.decrypt("")
        except Exception:
            raise ValueError("El PDF está protegido con contraseña") from None
    paginas = [(pagina.extract_text() or "").strip() for pagina in lector.pages]
    texto = "\n\n".join(p for p in paginas if p)
    if not texto.strip():
        raise ValueError("El PDF no tiene texto (¿es un escaneo?). Súbelo como texto o Word.")
    return texto


def _docx(contenido: bytes) -> str:
    from docx import Document
    documento = Document(io.BytesIO(contenido))
    lineas = []
    for parrafo in documento.paragraphs:
        texto = parrafo.text.strip()
        if not texto:
            continue
        estilo = (parrafo.style.name or "").lower() if parrafo.style is not None else ""
        lineas.append(f"# {texto}" if estilo.startswith(("heading", "título", "titulo", "title")) else texto)
    for tabla in documento.tables:
        for fila in tabla.rows:
            celdas = [c.text.strip() for c in fila.cells if c.text.strip()]
            if celdas:
                lineas.append(" | ".join(celdas))
    return "\n".join(lineas)


def _xlsx(contenido: bytes) -> str:
    from openpyxl import load_workbook
    libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
    partes = []
    for hoja in libro.worksheets:
        filas = []
        for fila in hoja.iter_rows(values_only=True):
            celdas = [str(c).strip() for c in fila if c is not None and str(c).strip()]
            if celdas:
                filas.append(" | ".join(celdas))
        if filas:
            partes.append(f"# {hoja.title}\n" + "\n".join(filas))
    libro.close()
    return "\n\n".join(partes)


def _csv(contenido: bytes) -> str:
    texto = _texto_plano(contenido)
    try:
        dialecto = csv.Sniffer().sniff(texto[:2000], delimiters=",;\t")
    except csv.Error:
        dialecto = csv.excel
    filas = [" | ".join(c.strip() for c in fila if c.strip()) for fila in csv.reader(io.StringIO(texto), dialecto)]
    return "\n".join(f for f in filas if f)


class _ExtractorHTML(HTMLParser):
    IGNORAR = {"script", "style", "noscript", "nav", "footer", "header", "svg", "form"}
    BLOQUES = {"p", "div", "li", "tr", "br", "section", "article", "table", "ul", "ol"}

    def __init__(self):
        super().__init__(convert_charrefs=True)
        self.partes = []
        self._ignorando = 0
        self._titulo = False

    def handle_starttag(self, tag, attrs):
        if tag in self.IGNORAR:
            self._ignorando += 1
        elif tag in ("h1", "h2", "h3", "h4"):
            self.partes.append("\n# ")
            self._titulo = True
        elif tag in self.BLOQUES:
            self.partes.append("\n")

    def handle_endtag(self, tag):
        if tag in self.IGNORAR and self._ignorando:
            self._ignorando -= 1
        elif tag in ("h1", "h2", "h3", "h4"):
            self.partes.append("\n")
            self._titulo = False
        elif tag in ("td", "th"):
            self.partes.append(" | ")

    def handle_data(self, data):
        if not self._ignorando:
            self.partes.append(" ".join(data.split()) + " ")


def texto_de_html(html: str) -> str:
    extractor = _ExtractorHTML()
    extractor.feed(html)
    lineas = [" ".join(l.split()).strip(" |") for l in "".join(extractor.partes).splitlines()]
    return "\n".join(l for l in lineas if l and l != "#")


LECTORES = {
    ".pdf": _pdf, ".docx": _docx, ".xlsx": _xlsx, ".csv": _csv,
    ".html": lambda c: texto_de_html(_texto_plano(c)), ".htm": lambda c: texto_de_html(_texto_plano(c)),
    ".txt": _texto_plano, ".md": _texto_plano,
}


def extraer_texto(nombre: str, contenido: bytes) -> str:
    """El texto del fichero según su extensión. ValueError si no se sabe leer o no trae texto."""
    extension = os.path.splitext((nombre or "").lower())[1] or ".txt"
    lector = LECTORES.get(extension)
    if lector is None:
        raise ValueError(f"No sé leer ficheros {extension}. Sube uno de estos: {', '.join(EXTENSIONES)}")
    try:
        texto = lector(contenido)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError(f"No se pudo leer el fichero ({type(exc).__name__}). ¿Está dañado?") from None
    texto = texto.strip()
    if not texto:
        raise ValueError("El documento está vacío")
    return texto[:MAXIMO_CARACTERES]


def _direccion_publica(url: str) -> None:
    """Solo webs públicas: nada de 127.0.0.1, la red de casa ni metadatos de la nube (SSRF)."""
    import ipaddress
    import socket
    from urllib.parse import urlsplit
    partes = urlsplit(url)
    if partes.scheme not in ("http", "https") or not partes.hostname:
        raise ValueError("La dirección tiene que empezar por http:// o https://")
    try:
        direcciones = {info[4][0] for info in socket.getaddrinfo(partes.hostname, partes.port or 443)}
    except socket.gaierror:
        raise ValueError("Esa web no existe o no se encuentra") from None
    for direccion in direcciones:
        ip = ipaddress.ip_address(direccion.split("%")[0])
        if not ip.is_global:
            raise ValueError("Solo se pueden leer webs públicas")


def leer_web(url: str, timeout: float = 20, maximo_bytes: int = 5 * 1024 * 1024) -> str:
    """El texto de una página web (la del negocio, su carta...). ValueError si no se puede."""
    import requests
    url = (url or "").strip()
    for _ in range(4):   # redirecciones, comprobando cada destino
        _direccion_publica(url)
        try:
            respuesta = requests.get(url, timeout=timeout, allow_redirects=False, stream=True,
                                     headers={"User-Agent": "femix/1.0 (+documentos del negocio)"})
        except requests.RequestException as exc:
            raise ValueError(f"No se pudo abrir la web ({type(exc).__name__})") from None
        if respuesta.is_redirect and respuesta.headers.get("location"):
            from urllib.parse import urljoin
            url = urljoin(url, respuesta.headers["location"])
            continue
        break
    else:
        raise ValueError("La web redirige demasiadas veces")
    if respuesta.status_code >= 400:
        raise ValueError(f"La web respondió con error {respuesta.status_code}")
    contenido = respuesta.raw.read(maximo_bytes + 1, decode_content=True)
    if len(contenido) > maximo_bytes:
        raise ValueError("La página es demasiado grande")
    tipo = respuesta.headers.get("content-type", "")
    return extraer_texto("web.pdf" if "pdf" in tipo else "web.html", contenido)
